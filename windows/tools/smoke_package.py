"""Exercise the frozen CLI against real conversion fixtures.

Run after packaging: .venv\Scripts\python.exe tools\smoke_package.py
Requires the development Pillow, PyMuPDF and python-docx dependencies for output
inspection. It uses temporary inputs and does not require LibreOffice/Tesseract.
"""
from __future__ import annotations

import argparse
import hashlib
import os
from pathlib import Path
import subprocess
import tempfile
import wave
import zipfile

from PIL import Image, ImageDraw, ImageOps
import pillow_heif
import pymupdf
from docx import Document

pillow_heif.register_heif_opener()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    default = Path(__file__).resolve().parents[2] / "dist/windows/FormatWheel/FormatWheel-cli.exe"
    parser.add_argument("--executable", type=Path, default=default)
    parser.add_argument("--ascii", action="store_true", help="Use ASCII names when diagnosing an older build")
    parser.add_argument("--skip-cutout", action="store_true", help="For packages intentionally built without a model")
    args = parser.parse_args()
    executable = args.executable.resolve()
    checks = 0
    if not executable.is_file():
        parser.error(f"Build the package first: {executable}")

    def invoke(*arguments):
        process = subprocess.run([str(executable), *map(str, arguments)], capture_output=True,
                                 encoding="utf-8", errors="replace", timeout=120,
                                 env=os.environ | {"QT_QPA_PLATFORM": "offscreen"})
        if process.returncode:
            raise AssertionError(f"Frozen CLI exit {process.returncode}: {' '.join(map(str, arguments))}\n{process.stdout}\n{process.stderr}")
        return process.stdout

    def run(action: str, *files: Path) -> list[Path]:
        nonlocal checks
        output = invoke("--run", action, *files)
        paths = [Path(line.strip()[3:]) for line in output.splitlines() if line.strip().startswith("-> ")]
        if not paths or not all(path.exists() for path in paths):
            raise AssertionError(f"No readable output paths from {action}: {output}")
        checks += 1
        print(f"PASS {action}")
        return paths

    ids = invoke("--ids").splitlines()
    assert len(ids) == len(set(ids)) == 73, "Frozen action catalog differs from upstream"
    checks += 1
    with tempfile.TemporaryDirectory(prefix="FormatWheel-package-smoke-") as temporary:
        folder = Path(temporary) / ("fixtures" if args.ascii else "中文 fixtures")
        folder.mkdir()
        corrupt_pdf = folder / "corrupt.pdf"
        corrupt_pdf.write_bytes(b"not a readable PDF")
        assert "corrupt.pdf:" in invoke("--scan", corrupt_pdf)
        checks += 1
        print("PASS malformed PDF privacy scan")
        source = folder / "transparent image.png"
        image = Image.new("RGBA", (81, 61), (200, 40, 70, 120))
        image.putpixel((0, 0), (0, 0, 0, 0))
        image.save(source)
        original_hash = hashlib.sha256(source.read_bytes()).digest()
        for fmt in ("png", "jpg", "heic", "webp", "avif", "gif", "tiff"):
            output = run("img.to." + fmt, source)[0]
            with Image.open(output) as result:
                result.load()
                assert result.size == (81, 61), fmt
                if fmt in {"jpg", "avif"}:
                    assert result.mode == "RGB", fmt
        output = run("img.half", source)[0]
        with Image.open(output) as result:
            assert result.size == (41, 31)
        output = run("img.rotate", source)[0]
        with Image.open(output) as result:
            assert result.size == (61, 81)
        run("img.mono", source)
        output = run("img.backdrop", source)[0]
        with Image.open(output) as result:
            assert result.width > 81 and result.height > 61
        for fmt, writer in (("jpg", "JPEG"), ("heic", "HEIF"), ("tiff", "TIFF")):
            camera = folder / ("private." + fmt)
            exif = Image.Exif()
            exif[315], exif[272], exif[274] = "Private author", "Private camera", 6
            Image.new("RGB", (40, 30), "red").save(camera, format=writer, exif=exif)
            with Image.open(camera) as before:
                expected_size = ImageOps.exif_transpose(before).size
            output = run("img.strip", camera)[0]
            with Image.open(output) as result:
                assert not result.getexif().get(315) and not result.getexif().get(272), fmt
                assert ImageOps.exif_transpose(result).size == expected_size, fmt
        pdf = folder / "two pages.pdf"
        with pymupdf.open() as doc:
            for index in range(2):
                page = doc.new_page(width=300, height=200)
                page.insert_text((20, 40), f"Heading {index + 1}", fontsize=24, color=(1, 0, 0))
                page.insert_text((20, 80), "Text remains editable.", fontsize=12)
            doc.set_metadata({"author": "Private author", "creator": "Private creator"})
            doc.save(pdf)
        for fmt in ("png", "jpg", "long", "docx", "txt", "md", "rtf", "html"):
            output = run("pdf.to." + fmt, pdf)[0]
            if fmt in {"png", "jpg"}:
                assert len(list(output.glob("*." + fmt))) == 2
            elif fmt == "docx":
                assert "Heading 1" in "\n".join(p.text for p in Document(output).paragraphs)
            elif fmt == "long":
                with Image.open(output) as result:
                    assert result.width == 1242 and result.height > 1600
        output = run("pdf.mono", pdf)[0]
        with pymupdf.open(output) as doc:
            assert "Heading 1" in doc[0].get_text()
            samples = doc[0].get_pixmap(colorspace=pymupdf.csRGB).samples
            assert samples[0::3] == samples[1::3] == samples[2::3]
        output = run("pdf.rotate", pdf)[0]
        with pymupdf.open(output) as doc:
            assert doc[0].rotation == 90
        output = run("pdf.strip", pdf)[0]
        with pymupdf.open(output) as doc:
            assert not doc.metadata.get("author") and not doc.metadata.get("creator")
            assert "Text remains editable" in doc[0].get_text()
        output = run("pdf.split", pdf)[0]
        assert len(list(output.glob("*.pdf"))) == 2
        output = run("pdf.merge", pdf, source)[0]
        with pymupdf.open(output) as doc:
            assert len(doc) == 3
        output = run("img.to.pdf", source)[0]
        with pymupdf.open(output) as doc:
            assert len(doc) == 1
        text = folder / "notes.txt"
        text.write_text("Example paragraph.\n" * 180, encoding="utf-8")
        for fmt in ("docx", "html", "pdf", "md"):
            output = run("doc.to." + fmt, text)[0]
            assert output.stat().st_size > 0
            if fmt == "pdf":
                with pymupdf.open(output) as doc:
                    assert len(doc) > 1 and "Example" in doc[0].get_text()
        word = folder / "document.docx"
        document = Document()
        document.core_properties.author = "Private author"
        document.add_heading("Document heading", level=1)
        p = document.add_paragraph()
        p.add_run("A normal ")
        p.add_run("bold phrase").bold = True
        document.save(word)
        run("doc.to.txt", word)
        output = run("doc.to.md", word)[0]
        assert "**bold phrase**" in output.read_text(encoding="utf-8")
        output = run("doc.strip", word)[0]
        assert not Document(output).core_properties.author
        output = run("gen.zip", source, pdf)[0]
        with zipfile.ZipFile(output) as archive:
            assert source.name in archive.namelist() and pdf.name in archive.namelist()
        nested = folder / "nested"
        nested.mkdir()
        inside = nested / "file.txt"
        inside.write_text("archive content", encoding="utf-8")
        output = run("gen.zip", inside, nested)[0]
        with zipfile.ZipFile(output) as archive:
            assert "nested/file.txt" in archive.namelist()
            assert not any(name.endswith(output.name) for name in archive.namelist())
        if not args.skip_cutout:
            subject = folder / "subject.png"
            image = Image.new("RGB", (320, 320), "white")
            ImageDraw.Draw(image).ellipse((60, 30, 260, 290), fill="navy")
            image.save(subject)
            output = run("img.cutout", subject)[0]
            with Image.open(output) as result:
                assert result.mode == "RGBA" and result.getchannel("A").getextrema() == (0, 255)
        for flag, name in (("--snapshot", "wheel.png"), ("--glyphs", "glyphs.png"), ("--weapons", "weapons.png")):
            output = folder / name
            invoke(flag, output)
            with Image.open(output) as result:
                result.load()
                assert min(result.size) > 100 and result.getbbox(), flag
            checks += 1
            print("PASS offscreen " + flag)
        sounds = folder / "sounds"
        invoke("--sounds", sounds)
        tracks = list(sounds.glob("*.wav"))
        assert len(tracks) >= 8
        for track in tracks:
            with wave.open(str(track)) as audio:
                assert audio.getnframes() > 0 and audio.getframerate() >= 22050
        checks += 1
        print("PASS sound exports")
        assert hashlib.sha256(source.read_bytes()).digest() == original_hash, "Input image was modified"
    print(f"Frozen package: {checks} checks passed.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
