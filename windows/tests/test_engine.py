"""Conversion fixtures check real pixels, PDF pages, metadata and unchanged inputs."""
from __future__ import annotations

from pathlib import Path
import hashlib
import os
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from PIL import Image, ImageChops, ImageOps
from formatwheel.engine import Engine, FormatWheelError, action_availability, scan_privacy

try:
    import pymupdf
except ImportError:
    pymupdf = None
try:
    from docx import Document
except ImportError:
    Document = None


class EngineFixtures(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="FormatWheel-tests-")
        self.folder = Path(self.temp.name) / "中文 files"
        self.folder.mkdir()
        self.png = self.folder / "transparent sample.png"
        image = Image.new("RGBA", (81, 61), (220, 30, 90, 80))
        image.putpixel((0, 0), (0, 0, 0, 0))
        image.save(self.png)
        self.engine = Engine()

    def tearDown(self):
        self.temp.cleanup()

    def test_half_rounds_up_and_retains_transparency(self):
        output = self.engine.run("img.half", [self.png]).outputs[0]
        with Image.open(output) as image:
            self.assertEqual(image.size, (41, 31))
            self.assertIn("A", image.getbands())

    def test_rotate_clockwise_matches_pixels(self):
        output = self.engine.run("img.rotate", [self.png]).outputs[0]
        with Image.open(self.png) as original, Image.open(output) as result:
            self.assertEqual(result.size, (61, 81))
            self.assertIsNone(ImageChops.difference(original.transpose(Image.Transpose.ROTATE_270), result).getbbox())

    def test_monochrome_keeps_alpha(self):
        output = self.engine.run("img.mono", [self.png]).outputs[0]
        with Image.open(output) as image:
            r, g, b, a = image.split()
            self.assertEqual(r.tobytes(), g.tobytes())
            self.assertEqual(r.tobytes(), b.tobytes())
            self.assertEqual(a.getextrema()[0], 0)

    def test_jpeg_transparency_flattens_white(self):
        output = self.engine.run("img.to.jpg", [self.png]).outputs[0]
        with Image.open(output) as image:
            self.assertEqual(image.mode, "RGB")
            self.assertGreater(min(image.getpixel((0, 0))), 200)

    def test_backdrop_has_padding_and_opaque_canvas(self):
        output = self.engine.run("img.backdrop", [self.png]).outputs[0]
        with Image.open(output) as image:
            self.assertGreater(image.width, 81)
            self.assertGreater(image.height, 61)
            self.assertEqual(image.mode, "RGB")

    def test_all_image_writers_decode(self):
        for fmt in ("png", "jpg", "heic", "webp", "avif", "gif", "tiff", "bmp"):
            with self.subTest(fmt=fmt):
                action = "img.to." + fmt
                if not action_availability(action, [self.png])[0]:
                    continue
                output = self.engine.run(action, [self.png]).outputs[0]
                with Image.open(output) as image:
                    image.load()
                    self.assertEqual(image.size, (81, 61))

    def test_outputs_never_overwrite_source_or_each_other(self):
        before = self.png.read_bytes()
        first = self.engine.run("img.to.png", [self.png]).outputs[0]
        second = self.engine.run("img.to.png", [self.png]).outputs[0]
        self.assertNotEqual(first, self.png)
        self.assertNotEqual(first, second)
        self.assertEqual(self.png.read_bytes(), before)

    def test_jpeg_strip_preserves_entropy_and_orientation(self):
        source = self.folder / "camera.jpg"
        exif = Image.Exif()
        exif[274], exif[272], exif[315] = 6, "Test camera", "Private author"
        Image.new("RGB", (80, 40), (120, 50, 70)).save(source, exif=exif)
        original = source.read_bytes()
        self.assertGreater(scan_privacy(source).level, 0)
        output = self.engine.run("img.strip", [source]).outputs[0]
        cleaned = output.read_bytes()
        self.assertEqual(original[original.index(b"\xff\xda"):], cleaned[cleaned.index(b"\xff\xda"):])
        self.assertNotIn(b"Private author", cleaned)
        with Image.open(output) as image:
            self.assertEqual(image.getexif().get(274), 6)
            self.assertEqual(ImageOps.exif_transpose(image).size, (40, 80))
        self.assertEqual(scan_privacy(output).level, 0)
        self.assertEqual(source.read_bytes(), original)

    def test_png_strip_preserves_pixels_and_orientation(self):
        source = self.folder / "camera.png"
        exif = Image.Exif()
        exif[274], exif[315] = 6, "Private author"
        Image.new("RGB", (80, 40), "red").save(source, exif=exif)
        output = self.engine.run("img.strip", [source]).outputs[0]
        with Image.open(source) as before, Image.open(output) as after:
            self.assertEqual(before.tobytes(), after.tobytes())
            self.assertEqual(after.getexif().get(274), 6)
            self.assertEqual(after.getexif().get(315), None)

    def test_cmyk_jpeg_strip_keeps_color_marker(self):
        source = self.folder / "print.jpg"
        Image.new("CMYK", (30, 20), (100, 20, 0, 30)).save(source)
        output = self.engine.run("img.strip", [source]).outputs[0]
        with Image.open(source) as before, Image.open(output) as after:
            self.assertEqual(before.convert("RGB").tobytes(), after.convert("RGB").tobytes())

    def test_heic_and_tiff_privacy_writers_do_not_inherit_exif(self):
        for ext, fmt in (("heic", "HEIF"), ("tiff", "TIFF")):
            if ext == "heic" and not action_availability("img.to.heic")[0]:
                continue
            with self.subTest(fmt=fmt):
                source = self.folder / ("private." + ext)
                exif = Image.Exif()
                exif[315], exif[272] = "Private author", "Camera model"
                Image.new("RGB", (30, 20), "red").save(source, format=fmt, exif=exif)
                output = self.engine.run("img.strip", [source]).outputs[0]
                self.assertEqual(scan_privacy(output).level, 0)

    def test_corrupt_cutout_model_rejected_before_import_or_network(self):
        from formatwheel.engine import _verified_model
        model = self.folder / "u2netp.onnx"
        model.write_bytes(b"corrupt model")
        with patch("formatwheel.engine._model_path", return_value=model):
            with self.assertRaisesRegex(FormatWheelError, "damaged"):
                _verified_model()

    def test_cutout_uses_local_model_with_network_disabled(self):
        from PIL import ImageDraw
        if not action_availability("img.cutout")[0]:
            self.skipTest("Local cutout model and runtime required")
        source = self.folder / "subject.png"
        image = Image.new("RGB", (320, 320), "white")
        ImageDraw.Draw(image).ellipse((60, 30, 260, 290), fill="navy")
        image.save(source)
        with patch("socket.socket.connect", side_effect=AssertionError("Unexpected network access")):
            output = self.engine.run("img.cutout", [source]).outputs[0]
        with Image.open(output) as result:
            self.assertEqual(result.mode, "RGBA")
            self.assertEqual(result.size, (320, 320))
            self.assertEqual(result.getchannel("A").getextrema(), (0, 255))

    def test_regular_image_transform_retains_author_metadata(self):
        source = self.folder / "camera.jpg"
        exif = Image.Exif()
        exif[274], exif[315] = 6, "Kept author"
        Image.new("RGB", (80, 40), "red").save(source, exif=exif)
        output = self.engine.run("img.to.png", [source]).outputs[0]
        with Image.open(output) as image:
            self.assertEqual(image.size, (40, 80))
            self.assertEqual(image.getexif().get(315), "Kept author")
            self.assertIn(image.getexif().get(274, 1), (1,))

    def test_compression_failure_cleans_reserved_output(self):
        source = self.folder / "tiny.jpg"
        Image.new("RGB", (1, 1), "white").save(source, quality=10, optimize=True)
        before = set(self.folder.iterdir())
        with self.assertRaises(FormatWheelError):
            self.engine.run("img.compress", [source])
        self.assertEqual(set(self.folder.iterdir()), before)

    def test_generic_hash_clipboard_and_paths(self):
        result = self.engine.run("gen.sha", [self.png])
        self.assertEqual(result.clip, hashlib.sha256(self.png.read_bytes()).hexdigest())
        self.assertEqual(self.engine.run("gen.path", [self.png]).clip, str(self.png.resolve()))
        self.assertEqual(self.engine.run("gen.name", [self.png]).clip, self.png.name)

    def test_archive_folders_duplicate_names_and_empty_directory(self):
        a, b = self.folder / "a", self.folder / "b"
        a.mkdir()
        b.mkdir()
        first, second = a / "same.txt", b / "same.txt"
        first.write_text("first")
        second.write_text("second")
        empty = a / "empty"
        empty.mkdir()
        output = self.engine.run("gen.zip", [first, second, empty]).outputs[0]
        with zipfile.ZipFile(output) as archive:
            self.assertEqual(archive.read("same.txt"), b"first")
            self.assertEqual(archive.read("same 2.txt"), b"second")
            self.assertIn("empty/", archive.namelist())

    def test_archive_does_not_include_itself_when_folder_is_also_selected(self):
        nested = self.folder / "nested"
        nested.mkdir()
        source = nested / "file.txt"
        source.write_text("contents", encoding="utf-8")
        output = self.engine.run("gen.zip", [source, nested]).outputs[0]
        with zipfile.ZipFile(output) as archive:
            self.assertIn("file.txt", archive.namelist())
            self.assertIn("nested/file.txt", archive.namelist())
            self.assertFalse(any(name.endswith(output.name) for name in archive.namelist()))

    def test_document_count_and_copy_chinese_english(self):
        source = self.folder / "sample.txt"
        source.write_text("中文 Hello world 123!", encoding="utf-8")
        result = self.engine.run("doc.count", [source])
        self.assertIn("5 words/Chinese characters", result.note)
        self.assertIn("16 characters", result.note)
        self.assertEqual(self.engine.run("doc.copy", [source]).clip, "中文 Hello world 123!")

    def test_markdown_plain_matches_upstream(self):
        source = self.folder / "notes.md"
        source.write_text("# Heading\n**Bold** `code`", encoding="utf-8")
        result = self.engine.run("doc.to.txt", [source])
        self.assertEqual(result.outputs[0].read_text(encoding="utf-8"), "Heading\nBold code")

    def test_majority_batch_skips_unrelated_file_and_progress_finishes(self):
        second = self.folder / "other.png"
        Image.new("RGB", (20, 10), "red").save(second)
        text = self.folder / "note.txt"
        text.write_text("skip")
        values = []
        result = self.engine.run("img.half", [self.png, text, second], values.append)
        self.assertEqual(len(result.outputs), 2)
        self.assertEqual(values[-1], 1)
        self.assertTrue(all(0 <= value <= 1 for value in values))

    def test_partial_failure_reports_successful_results(self):
        broken = self.folder / "broken.png"
        broken.write_bytes(b"not image data")
        result = self.engine.run("img.half", [self.png, broken])
        self.assertEqual(len(result.outputs), 1)
        self.assertIn("1 failed", result.note)

    def test_missing_dependency_message_is_actionable(self):
        with patch("formatwheel.engine.find_tool", return_value=None):
            available, reason = action_availability("vid.to.mp4", [self.png])
            self.assertFalse(available)
            self.assertIn("FFmpeg", reason)
            with self.assertRaisesRegex(FormatWheelError, "FFmpeg"):
                self.engine.run("vid.to.mp4", [self.png])

    def test_unknown_action_leaves_no_output(self):
        before = set(self.folder.iterdir())
        with self.assertRaisesRegex(FormatWheelError, "Unknown action"):
            self.engine.run("img.explode", [self.png])
        self.assertEqual(set(self.folder.iterdir()), before)

    def test_corrupt_metadata_does_not_block_the_wheel(self):
        for ext in ("pdf", "docx", "png"):
            source = self.folder / ("corrupt." + ext)
            source.write_bytes(b"not a readable document")
            self.assertEqual(scan_privacy(source).level, 0)


@unittest.skipUnless(pymupdf is not None, "PyMuPDF required")
class PDFEngineFixtures(unittest.TestCase):
    setUp = EngineFixtures.setUp
    tearDown = EngineFixtures.tearDown
    def make_pdf(self, name="sample.pdf", pages=2):
        path = self.folder / name
        with pymupdf.open() as doc:
            for index in range(pages):
                page = doc.new_page(width=300, height=200)
                page.insert_text((20, 40), f"Heading {index + 1}", fontsize=24, color=(1, 0, 0))
                page.insert_text((20, 80), "Body text can be extracted.", fontsize=12)
            doc.set_metadata({"author": "Secret author", "creator": "Private software", "subject": "Secret subject"})
            doc.save(path)
        return path

    def test_pdf_split_merge_rotation_and_raster(self):
        source = self.make_pdf()
        original = source.read_bytes()
        split = Engine.run("pdf.split", [source]).outputs[0]
        self.assertEqual(len(list(split.glob("*.pdf"))), 2)
        merged = Engine.run("pdf.merge", [*sorted(split.glob("*.pdf")), self.png]).outputs[0]
        with pymupdf.open(merged) as doc:
            self.assertEqual(len(doc), 3)
            self.assertIn("Heading 1", doc[0].get_text())
            self.assertIn("PyMuPDF", doc.metadata["producer"])
        rotated = Engine.run("pdf.rotate", [source]).outputs[0]
        with pymupdf.open(rotated) as doc:
            self.assertEqual(doc[0].rotation, 90)
        folder = Engine.run("pdf.to.png", [rotated]).outputs[0]
        with Image.open(sorted(folder.glob("*.png"))[0]) as image:
            self.assertEqual(image.size, (400, 600))
        self.assertEqual(source.read_bytes(), original)

    def test_pdf_monochrome_keeps_text_layer(self):
        source = self.make_pdf()
        result = Engine.run("pdf.mono", [source]).outputs[0]
        with pymupdf.open(result) as doc:
            self.assertIn("Heading 1", doc[0].get_text())
            pix = doc[0].get_pixmap(colorspace=pymupdf.csRGB)
            data = pix.samples
            self.assertEqual(data[0::3], data[1::3])
            self.assertEqual(data[0::3], data[2::3])

    def test_pdf_strip_removes_personal_metadata(self):
        source = self.make_pdf()
        self.assertGreater(scan_privacy(source).level, 0)
        result = Engine.run("pdf.strip", [source]).outputs[0]
        self.assertEqual(scan_privacy(result).level, 0)
        self.assertNotIn(b"Secret author", result.read_bytes())
        with pymupdf.open(result) as doc:
            self.assertIn("Body text", doc[0].get_text())
            self.assertIn("PyMuPDF", doc.metadata["producer"])
            self.assertIn("Artifex", doc.metadata["producer"])

    def test_pdf_long_image_and_text_outputs(self):
        source = self.make_pdf()
        output = Engine.run("pdf.to.long", [source]).outputs[0]
        with Image.open(output) as image:
            self.assertEqual(image.width, 1242)
            self.assertGreater(image.height, 1600)
        for fmt in ("txt", "md", "rtf", "html", "docx"):
            if fmt == "docx" and Document is None:
                continue
            output = Engine.run("pdf.to." + fmt, [source]).outputs[0]
            self.assertGreater(output.stat().st_size, 0)
            if fmt == "txt":
                self.assertIn("Heading 1", output.read_text(encoding="utf-8"))

    def test_pdf_password_rejected(self):
        source = self.folder / "locked.pdf"
        with pymupdf.open() as doc:
            doc.new_page()
            doc.save(source, encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="private")
        with self.assertRaisesRegex(FormatWheelError, "password"):
            Engine.run("pdf.to.png", [source])

    def test_one_page_split_rejected(self):
        with self.assertRaisesRegex(FormatWheelError, "one page"):
            Engine.run("pdf.split", [self.make_pdf(pages=1)])

    def test_plain_text_pdf_has_unicode_and_multiple_pages(self):
        source = self.folder / "notes.txt"
        source.write_text("中文 Example paragraph.\n" * 180, encoding="utf-8")
        output = Engine.run("doc.to.pdf", [source]).outputs[0]
        with pymupdf.open(output) as doc:
            self.assertGreater(len(doc), 1)
            self.assertIn("PyMuPDF", doc.metadata["producer"])
            self.assertIn("Example", "".join(p.get_text() for p in doc))
            self.assertIn("中文", "".join(p.get_text() for p in doc))


@unittest.skipUnless(Document is not None, "python-docx required")
class WordFixtures(unittest.TestCase):
    setUp = EngineFixtures.setUp
    tearDown = EngineFixtures.tearDown

    def test_word_markdown_preserves_inline_bold(self):
        source = self.folder / "emphasis.docx"
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run("A normal ")
        paragraph.add_run("bold phrase").bold = True
        paragraph.add_run(" in this sentence.")
        document.save(source)
        output = Engine.run("doc.to.md", [source]).outputs[0]
        self.assertIn("A normal **bold phrase** in this sentence.", output.read_text(encoding="utf-8"))
    def test_word_strip_and_extract_preserves_document_parts(self):
        source = self.folder / "sample.docx"
        document = Document()
        document.core_properties.author = "Secret author"
        document.core_properties.last_modified_by = "Secret editor"
        document.add_heading("A heading", level=1)
        document.add_paragraph("Hello 中文 content.")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "Table text"
        document.save(source)
        self.assertGreaterEqual(scan_privacy(source).level, 3)
        output = Engine.run("doc.strip", [source]).outputs[0]
        self.assertEqual(scan_privacy(output).level, 0)
        with zipfile.ZipFile(source) as original, zipfile.ZipFile(output) as cleaned:
            self.assertEqual(original.read("word/document.xml"), cleaned.read("word/document.xml"))
        text = Engine.run("doc.copy", [source]).clip
        self.assertIn("Table text", text)
        self.assertIn("中文", text)
        markdown = Engine.run("doc.to.md", [source]).outputs[0].read_text(encoding="utf-8")
        self.assertIn("# A heading", markdown)


if __name__ == "__main__":
    unittest.main()
