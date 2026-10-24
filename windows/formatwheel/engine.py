"""Local Windows implementations of every FormatWheel action.

The engine never mutates an input or the clipboard. Commands use argument arrays,
outputs are reserved exclusively, and failed writes are removed.
"""
from __future__ import annotations

from collections import Counter
from contextlib import contextmanager
from dataclasses import dataclass, field
from html import escape
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path
from typing import Callable, Iterable
import hashlib
import importlib.util
import json
import math
import os
import plistlib
import re
import shutil
import stat
import struct
import subprocess
import sys
import tempfile
import threading
import time
import unicodedata
import xml.etree.ElementTree as ET
import zipfile
import zlib

from PIL import Image, ImageDraw, ImageFilter, ImageOps

try:
    import pillow_heif
    pillow_heif.register_heif_opener()
except ImportError:
    pillow_heif = None

class FormatWheelError(RuntimeError):
    """An actionable conversion error suitable for the result banner."""


class OperationCancelled(FormatWheelError):
    """A user cancellation; partial results have been removed."""


class OperationControl:
    """Cooperative controls shared between a worker and its user interface."""

    def __init__(self):
        self._condition = threading.Condition()
        self._paused = False
        self._cancelled = False
        self._paused_since: float | None = None
        self._paused_total = 0.0

    @property
    def is_paused(self) -> bool:
        with self._condition:
            return self._paused

    @property
    def is_cancelled(self) -> bool:
        with self._condition:
            return self._cancelled

    @property
    def paused_seconds(self) -> float:
        with self._condition:
            return self._paused_total + (time.monotonic() - self._paused_since if self._paused_since is not None else 0)

    def pause(self):
        with self._condition:
            if not self._paused and not self._cancelled:
                self._paused = True
                self._paused_since = time.monotonic()

    def _end_pause(self):
        if self._paused_since is not None:
            self._paused_total += time.monotonic() - self._paused_since
            self._paused_since = None
        self._paused = False

    def resume(self):
        with self._condition:
            self._end_pause()
            self._condition.notify_all()

    def cancel(self):
        with self._condition:
            self._cancelled = True
            self._end_pause()
            self._condition.notify_all()

    def checkpoint(self, on_pause: Callable[[], None] | None = None):
        announced = False
        while True:
            with self._condition:
                if self._cancelled:
                    raise OperationCancelled("Operation cancelled. The partial output was removed.")
                if not self._paused:
                    return
            if not announced and on_pause:
                # Never invoke external callbacks while holding the condition lock.
                announced = True
                on_pause()
            with self._condition:
                if self._paused and not self._cancelled:
                    self._condition.wait()


@dataclass(frozen=True)
class ProgressUpdate:
    fraction: float | None = None
    phase: str = "scanning"
    processed_bytes: int = 0
    total_bytes: int = 0
    files_completed: int = 0
    files_total: int = 0
    current_file: str = ""
    output_path: Path | None = None
    archive_bytes: int = 0
    elapsed: float = 0.0
    speed_bps: float | None = None
    eta_seconds: float | None = None


Progress = Callable[[float | ProgressUpdate], None]


@dataclass
class RunResult:
    outputs: list[Path] = field(default_factory=list)
    note: str = ""
    clip: str | None = None


@dataclass
class Wanted:
    level: int = 0
    reasons: list[str] = field(default_factory=list)


def find_tool(name: str) -> str | None:
    """Resolve tools in portable distribution, explicit override, PATH or install locations."""
    override = os.environ.get("FORMATWHEEL_" + name.upper())
    if override and Path(override).is_file():
        return override
    bases = [Path(sys.executable).parent, Path(__file__).resolve().parents[1]]
    if getattr(sys, "_MEIPASS", None):
        bases.insert(0, Path(sys._MEIPASS))
    aliases = [name, "libreoffice"] if name == "soffice" else [name]
    for base in bases:
        for alias in aliases:
            for relative in (f"tools/{alias}.exe", f"tools/{alias}/{alias}.exe", f"{alias}.exe"):
                candidate = base / relative
                if candidate.is_file():
                    return str(candidate)
    for alias in aliases:
        found = shutil.which(alias)
        if found:
            return found
    program_dirs = [os.environ.get("ProgramFiles", "C:/Program Files"), os.environ.get("ProgramFiles(x86)", "C:/Program Files (x86)"), str(Path(os.environ.get("LOCALAPPDATA", str(Path.home() / "AppData/Local"))) / "Programs")]
    locations = {"soffice": ["LibreOffice/program/soffice.exe"], "tesseract": ["Tesseract-OCR/tesseract.exe"]}
    for base in program_dirs:
        for relative in locations.get(name, []):
            candidate = Path(base) / relative
            if candidate.is_file():
                return str(candidate)
    return None


def _require(name: str) -> str:
    found = find_tool(name)
    if not found:
        labels = {"ffmpeg": "FFmpeg", "ffprobe": "FFmpeg (ffprobe)", "soffice": "LibreOffice", "tesseract": "Tesseract OCR"}
        raise FormatWheelError(f"Install {labels.get(name, name)} using Install-Dependencies.ps1, then restart FormatWheel.")
    return found


def _command(args: list[str], timeout: int = 600) -> str:
    try:
        result = subprocess.run(args, stdin=subprocess.DEVNULL, capture_output=True,
                                text=True, encoding="utf-8", errors="replace", timeout=timeout,
                                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    except (OSError, subprocess.TimeoutExpired) as exc:
        raise FormatWheelError(f"{Path(args[0]).name}: {exc}") from exc
    if result.returncode:
        lines = result.stderr.strip().splitlines() or result.stdout.strip().splitlines()
        raise FormatWheelError(f"{Path(args[0]).name}: {lines[-1][:350] if lines else 'conversion failed'}")
    return result.stdout


def _kind(path: Path) -> str:
    from .catalog import file_kind
    return file_kind(path)


def _canonical(ext: str) -> str:
    return {"jpeg": "jpg", "jpe": "jpg", "tif": "tiff", "heif": "heic", "aif": "aiff", "aifc": "aiff", "oga": "opus", "ogg": "opus", "markdown": "md", "htm": "html", "text": "txt"}.get(ext.lower().lstrip("."), ext.lower().lstrip("."))


def _output_dir(source: Path) -> Path:
    for folder in (source.parent, Path.home() / "Downloads", Path(tempfile.gettempdir())):
        try:
            folder.mkdir(parents=True, exist_ok=True)
            with tempfile.TemporaryFile(dir=folder):
                pass
            return folder
        except OSError:
            continue
    raise FormatWheelError("No writable output folder is available.")


def _reserve(source: Path, ext: str, suffix: str = "", directory: bool = False) -> Path:
    folder = _output_dir(source)
    base = source.stem + (" " + suffix if suffix else "")
    for n in range(1, 100000):
        name = base + (f" {n}" if n > 1 else "") + ("." + ext if ext else "")
        target = folder / name
        try:
            if directory:
                target.mkdir()
            else:
                with target.open("xb"):
                    pass
            return target
        except FileExistsError:
            continue
    raise FormatWheelError("Too many files with the same output name.")


@contextmanager
def _output(source: Path, ext: str, suffix: str = ""):
    target = _reserve(source, ext, suffix)
    try:
        yield target
        if not target.is_file() or not target.stat().st_size:
            raise FormatWheelError("The converter did not produce a file.")
    except BaseException:
        target.unlink(missing_ok=True)
        raise


@contextmanager
def _folder(source: Path, suffix: str):
    target = _reserve(source, "", suffix, directory=True)
    try:
        yield target
    except BaseException:
        shutil.rmtree(target, ignore_errors=True)
        raise


def _fitz():
    try:
        import pymupdf
        return pymupdf
    except ImportError as exc:
        raise FormatWheelError("PyMuPDF is missing. Reinstall the Windows package.") from exc


def _open_pdf(path: Path):
    doc = _fitz().open(path)
    if doc.needs_pass:
        doc.close()
        raise FormatWheelError("This PDF is password protected.")
    if not len(doc):
        doc.close()
        raise FormatWheelError("This PDF has no pages.")
    return doc


def _stamp_pdf(doc):
    metadata = doc.metadata or {}
    metadata["producer"] = f"PyMuPDF {_fitz().VersionBind} / MuPDF (Artifex)"
    doc.set_metadata(metadata)


def _save_pdf(doc, path: Path, **options):
    # Preserve the AGPL library attribution independently of personal metadata.
    _stamp_pdf(doc)
    doc.save(path, **options)


def _stamp_pdf_file(path: Path):
    with _fitz().open(path) as doc:
        _stamp_pdf(doc)
        doc.saveIncr()


def _has_module(name: str) -> bool:
    try:
        return importlib.util.find_spec(name) is not None
    except (ImportError, ValueError):
        return False


def _model_path() -> Path:
    folders = [Path(os.environ.get("U2NET_HOME", str(Path.home() / ".u2net"))),
               Path(os.environ.get("LOCALAPPDATA", str(Path.home() / "AppData/Local"))) / "FormatWheel/models",
               Path(sys.executable).parent / "models",
               Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parents[1])) / "models",
               Path(__file__).resolve().parents[1] / ".tools/models"]
    for folder in folders:
        for name in ("u2netp.onnx", "u2net.onnx"):
            if (folder / name).is_file():
                return folder / name
    return folders[1] / "u2netp.onnx"


def _verified_model() -> Path:
    model = _model_path()
    if not model.is_file():
        raise FormatWheelError("Install the background-removal model with Install-Dependencies.ps1 -WithCutout.")
    expected = {"u2netp": "8e83ca70e441ab06c318d82300c84806", "u2net": "60024c5c889badc19c04ad937298a77b"}
    digest = hashlib.md5(usedforsecurity=False)
    with model.open("rb") as handle:
        for chunk in iter(lambda: handle.read(4 << 20), b""):
            digest.update(chunk)
    if digest.hexdigest() != expected.get(model.stem):
        raise FormatWheelError("The local background-removal model is damaged. Run Install-Dependencies.ps1 -WithCutout to repair it.")
    return model


def _cutout(image: Image.Image) -> Image.Image:
    model = _verified_model()
    import onnxruntime
    from rembg import remove
    from rembg.sessions.u2net import U2netSession
    from rembg.sessions.u2netp import U2netpSession
    base = U2netpSession if model.stem == "u2netp" else U2netSession
    class LocalSession(base):
        @classmethod
        def download_models(cls, *args, **kwargs):
            # Deliberately bypass rembg's downloader. Inference never accesses the network,
            # including when the file disappears between validation and session creation.
            return str(model)
    options = onnxruntime.SessionOptions()
    options.intra_op_num_threads = min(4, os.cpu_count() or 1)
    session = LocalSession(model.stem, options, providers=["CPUExecutionProvider"])
    return remove(image, session=session)


def action_availability(action_id: str, files: Iterable[Path] = ()) -> tuple[bool, str]:
    """Gate optional dependencies without hiding any upstream actions."""
    paths = [Path(p) for p in files]
    if action_id.startswith(("vid.", "aud.")):
        if not find_tool("ffmpeg") or not find_tool("ffprobe"):
            return False, "Install FFmpeg using Install-Dependencies.ps1."
    if action_id in ("img.ocr", "pdf.ocr") and not find_tool("tesseract"):
        return False, "Install Tesseract OCR and its English / Chinese language packs."
    if action_id == "img.cutout":
        if not _has_module("rembg") or not _model_path().is_file():
            return False, "Run Install-Dependencies.ps1 -WithCutout to install the local background-removal model."
    if action_id == "img.to.heic":
        try:
            if pillow_heif is None or not pillow_heif.libheif_info().get("HEIF"):
                return False, "This installation has no HEIC encoder. Reinstall pillow-heif."
        except Exception:
            return False, "This installation has no HEIC encoder."
    if action_id == "img.to.avif":
        Image.init()
        if "AVIF" not in Image.SAVE and not find_tool("ffmpeg"):
            return False, "Install FFmpeg or a Pillow build with AVIF support."
    if action_id == "doc.strip" and any(p.suffix.lower() != ".docx" for p in paths):
        return False, "Privacy stripping is available for .docx documents."
    if action_id == "gen.sha" and any(p.is_dir() for p in paths):
        return False, "SHA-256 is available for files, not folders."
    if action_id == "pdf.merge" and len(paths) < 2:
        return False, "Select at least two PDFs or images to merge."
    if action_id.startswith("doc.") and action_id != "doc.strip":
        target = action_id.removeprefix("doc.to.") if action_id.startswith("doc.to.") else None
        plain_sources = {"txt", "md", "docx", "html"}
        need_office = any(_canonical(p.suffix) not in plain_sources for p in paths)
        need_office |= target in {"doc", "rtf", "odt"}
        need_office |= target == "md" and any(_canonical(p.suffix) not in {"txt", "md", "docx"} for p in paths)
        need_office |= target in {"docx", "html", "pdf"} and any(_canonical(p.suffix) not in {"txt", "docx"} for p in paths)
        need_office |= target == "pdf" and any(_canonical(p.suffix) == "docx" for p in paths)
        if need_office and not find_tool("soffice"):
            return False, "Install LibreOffice using Install-Dependencies.ps1 for this document conversion."
    return True, ""


class Engine:
    @staticmethod
    def run(action_id: str, files: Iterable[Path | str], progress: Progress | None = None,
            control: OperationControl | None = None) -> RunResult:
        paths = [Path(p).expanduser().resolve() for p in files]
        if not paths:
            raise FormatWheelError("Drop at least one file first.")
        if any(not p.exists() for p in paths):
            raise FormatWheelError("An input file no longer exists.")
        callback = progress or (lambda _value: None)
        if action_id != "gen.zip":
            if control:
                control.checkpoint()
            callback(0)
        dominant = Counter(_kind(p) for p in paths).most_common(1)[0][0]
        targets = [p for p in paths if action_id.startswith("gen.") or _kind(p) == dominant]
        whole_selection = action_id in {"pdf.merge", "gen.zip", "gen.path", "gen.name"} or (action_id == "img.to.pdf" and len(paths) > 1)
        available, reason = action_availability(action_id, paths if whole_selection else targets)
        if not available:
            raise FormatWheelError(reason)
        if action_id in {"gen.path", "gen.name"}:
            text = "\n".join(str(p) if action_id == "gen.path" else p.name for p in paths)
            callback(1)
            return RunResult(note="Copied to the clipboard.", clip=text)
        if action_id == "gen.zip":
            return _zip(paths, callback, control or OperationControl())
        if action_id == "pdf.merge" or (action_id == "img.to.pdf" and len(paths) > 1):
            result = _merge(paths, callback)
            callback(1)
            return result
        result = RunResult()
        notes, clips, failures = [], [], []
        for index, path in enumerate(targets):
            step = lambda fraction, i=index: callback((i + max(0, min(1, fraction))) / len(targets))
            try:
                part = _run_one(action_id, path, step)
                result.outputs.extend(part.outputs)
                if part.note:
                    notes.append(part.note)
                if part.clip is not None:
                    clips.append(part.clip)
            except Exception as exc:
                failures.append(f"{path.name}: {exc}")
            step(1)
        if not result.outputs and not notes and not clips:
            raise FormatWheelError(failures[0] if failures else "No compatible input files.")
        result.note = " · ".join(notes if len(targets) == 1 else [f"{len(targets) - len(failures)} file(s) processed."])
        if failures:
            result.note += f" · {len(failures)} failed: {failures[0]}"
        result.clip = "\n\n".join(clips) if clips else None
        callback(1)
        return result


def _run_one(action: str, source: Path, progress: Progress) -> RunResult:
    if action.startswith("img."):
        return _image_action(action, source)
    if action.startswith("pdf."):
        return _pdf_action(action, source, progress)
    if action.startswith(("vid.", "aud.")):
        return _media_action(action, source, progress)
    if action.startswith("doc."):
        return _doc_action(action, source)
    if action == "gen.sha":
        digest = hashlib.sha256()
        with source.open("rb") as handle:
            for chunk in iter(lambda: handle.read(4 << 20), b""):
                digest.update(chunk)
        return RunResult(note=f"SHA-256 {digest.hexdigest()[:12]}… copied.", clip=digest.hexdigest())
    raise FormatWheelError(f"Unknown action: {action}")


_ZIP_CHUNK_SIZE = 1024 * 1024
_ZIP_PROGRESS_INTERVAL = 0.08


@dataclass(frozen=True)
class _ZipEntry:
    path: Path
    name: str
    size: int
    mtime_ns: int
    ctime_ns: int
    device: int
    inode: int
    directory: bool

    def unchanged(self, current: os.stat_result) -> bool:
        return (not self.directory and stat.S_ISREG(current.st_mode)
                and (current.st_size, current.st_mtime_ns, current.st_ctime_ns, current.st_dev, current.st_ino)
                == (self.size, self.mtime_ns, self.ctime_ns, self.device, self.inode))


class _ZipProgress:
    def __init__(self, callback: Callable[[ProgressUpdate], None], control: OperationControl):
        self.callback, self.control = callback, control
        self.started = time.monotonic()
        self.last_emitted = -float("inf")
        self.phase = "scanning"
        self.processed_bytes = self.total_bytes = 0
        self.files_completed = self.files_total = 0
        self.current_file = ""
        self.output_path: Path | None = None
        self.compression_started: float | None = None
        self.paused_before_compression = 0.0

    def emit(self, phase: str | None = None, *, force: bool = False):
        now = time.monotonic()
        if phase is not None and phase != self.phase:
            self.phase, force = phase, True
        if not force and now - self.last_emitted < _ZIP_PROGRESS_INTERVAL:
            return
        self.last_emitted = now
        if self.compression_started is None and self.phase == "compressing":
            self.compression_started = now
            self.paused_before_compression = self.control.paused_seconds
        if self.phase == "completed":
            fraction = 1.0
        elif self.phase == "scanning" or self.phase == "paused" and self.compression_started is None:
            fraction = None
        else:
            fraction = min(self.processed_bytes / self.total_bytes, math.nextafter(1.0, 0.0)) if self.total_bytes else 0.0
        speed = eta = None
        if self.phase == "compressing" and self.compression_started is not None:
            active = now - self.compression_started - (self.control.paused_seconds - self.paused_before_compression)
            if active >= .1 and self.processed_bytes:
                speed = self.processed_bytes / active
                if self.total_bytes > self.processed_bytes:
                    eta = (self.total_bytes - self.processed_bytes) / speed
        archive_bytes = 0
        if self.output_path is not None:
            try:
                archive_bytes = self.output_path.stat().st_size
            except OSError:
                pass
        self.callback(ProgressUpdate(fraction, self.phase, self.processed_bytes, self.total_bytes,
                                     self.files_completed, self.files_total, self.current_file,
                                     self.output_path, archive_bytes, now - self.started, speed, eta))

    def checkpoint(self):
        phase = self.phase
        paused = False
        def announce_pause():
            nonlocal paused
            paused = True
            self.emit("paused", force=True)
        self.control.checkpoint(announce_pause)
        if paused:
            self.emit(phase, force=True)


def _zip_link(path: Path, attributes: os.stat_result) -> bool:
    # Directory junctions need the same treatment as symlinks; OneDrive cloud
    # placeholders use other reparse tags and remain ordinary input files.
    return stat.S_ISLNK(attributes.st_mode) or getattr(attributes, "st_reparse_tag", 0) == getattr(stat, "IO_REPARSE_TAG_MOUNT_POINT", -1)


def _zip_scan(paths: list[Path], reporter: _ZipProgress) -> list[_ZipEntry]:
    entries: list[_ZipEntry] = []
    used: set[str] = set()

    def visit(path: Path, name: str):
        reporter.checkpoint()
        attributes = path.lstat()
        if _zip_link(path, attributes):
            return
        directory = stat.S_ISDIR(attributes.st_mode)
        if not directory and not stat.S_ISREG(attributes.st_mode):
            raise FormatWheelError(f"Cannot archive this special file: {path.name}")
        entry = _ZipEntry(path, name + "/" if directory else name,
                          0 if directory else attributes.st_size,
                          attributes.st_mtime_ns, attributes.st_ctime_ns,
                          attributes.st_dev, attributes.st_ino, directory)
        entries.append(entry)
        if not directory:
            reporter.total_bytes += entry.size
            reporter.files_total += 1
        reporter.current_file = name
        reporter.emit()
        if directory:
            with os.scandir(path) as children:
                ordered = []
                for child in children:
                    reporter.checkpoint()
                    ordered.append(child)
                    reporter.emit()
                ordered.sort(key=lambda child: child.name.casefold())
            for child in ordered:
                visit(Path(child.path), name + "/" + child.name)

    for path in paths:
        name, n = path.name, 2
        while name.casefold() in used:
            name = f"{path.stem} {n}{path.suffix}"
            n += 1
        used.add(name.casefold())
        try:
            visit(path, name)
        except OSError as exc:
            raise FormatWheelError(f"Could not scan {path.name}: {exc}") from exc
    return entries


def _zip_changed(entry: _ZipEntry):
    raise FormatWheelError(f"{entry.path.name} changed while creating the archive. No archive was kept; retry when the file has finished saving.")


def _zip(paths: list[Path], progress: Callable[[ProgressUpdate], None], control: OperationControl) -> RunResult:
    reporter = _ZipProgress(progress, control)
    reporter.emit("scanning", force=True)
    entries = _zip_scan(paths, reporter)
    reporter.checkpoint()
    first_chunk = True
    with _output(paths[0], "zip", "" if len(paths) == 1 else f"Archive {len(paths)} items") as out:
        reporter.output_path = out
        reporter.current_file = ""
        reporter.emit("compressing", force=True)
        reporter.checkpoint()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED, allowZip64=True) as archive:
            for entry in entries:
                reporter.checkpoint()
                reporter.current_file = entry.name
                # The complete manifest is made before output creation, so even a
                # selected ancestor directory can never contain this new archive.
                if entry.path.resolve() == out.resolve():
                    continue
                if entry.directory:
                    archive.writestr(zipfile.ZipInfo.from_file(entry.path, entry.name), b"")
                    reporter.emit()
                    continue
                try:
                    with entry.path.open("rb") as source:
                        if not entry.unchanged(os.fstat(source.fileno())):
                            _zip_changed(entry)
                        info = zipfile.ZipInfo.from_file(entry.path, entry.name)
                        info.compress_type = zipfile.ZIP_DEFLATED
                        info._compresslevel = archive.compresslevel
                        zip64 = entry.size * 1.05 > zipfile.ZIP64_LIMIT
                        remaining = entry.size
                        with archive.open(info, "w", force_zip64=zip64) as compressed:
                            while remaining:
                                reporter.checkpoint()
                                chunk = source.read(min(_ZIP_CHUNK_SIZE, remaining))
                                if not chunk:
                                    _zip_changed(entry)
                                compressed.write(chunk)
                                remaining -= len(chunk)
                                reporter.processed_bytes += len(chunk)
                                reporter.emit(force=first_chunk)
                                first_chunk = False
                            reporter.checkpoint()
                            if source.read(1) or not entry.unchanged(os.fstat(source.fileno())):
                                _zip_changed(entry)
                    reporter.files_completed += 1
                    reporter.emit()
                except OSError as exc:
                    raise FormatWheelError(f"Could not read {entry.path.name}: {exc}") from exc
            # Files processed earlier may have changed while later ones were read.
            # Validation can take time on large selections or network drives, so
            # show finalization before starting it and keep its details responsive.
            reporter.emit("finalizing", force=True)
            reporter.checkpoint()
            for entry in entries:
                reporter.checkpoint()
                if not entry.directory:
                    reporter.current_file = entry.name
                    try:
                        current = entry.path.stat()
                    except OSError:
                        _zip_changed(entry)
                    if not entry.unchanged(current):
                        _zip_changed(entry)
                    reporter.emit()
            reporter.checkpoint()
        reporter.checkpoint()
    reporter.current_file = ""
    reporter.emit("completed", force=True)
    return RunResult([out])


def _load_image(source: Path) -> tuple[Image.Image, dict]:
    with Image.open(source) as original:
        info = dict(original.info)
        upright = ImageOps.exif_transpose(original)
        exif = upright.getexif()
        for key in (274, 40962, 40963):
            if key in exif:
                del exif[key]
        info["exif"] = exif.tobytes() if exif else b""
        return upright.copy(), info


def _alpha(image: Image.Image) -> bool:
    return image.mode in ("RGBA", "LA") and image.getchannel("A").getextrema()[0] < 250 or image.mode == "P" and "transparency" in image.info


def _flatten(image: Image.Image) -> Image.Image:
    rgba = image.convert("RGBA")
    canvas = Image.new("RGB", image.size, "white")
    canvas.paste(rgba, mask=rgba.getchannel("A"))
    return canvas


def _save_image(image: Image.Image, output: Path, fmt: str, quality: int | None = None, metadata: dict | None = None):
    formats = {"jpg": "JPEG", "png": "PNG", "heic": "HEIF", "tiff": "TIFF", "gif": "GIF", "bmp": "BMP", "webp": "WEBP", "avif": "AVIF"}
    if fmt not in formats:
        raise FormatWheelError(f"Unsupported image format: {fmt}")
    # HEIF writers inherit image.info automatically. Supply only explicit metadata,
    # otherwise a privacy export could silently carry the source's EXIF/XMP.
    image.info = {"transparency": image.info["transparency"]} if "transparency" in image.info else {}
    if fmt in {"jpg", "bmp", "avif"}:
        image = _flatten(image)
    elif image.mode not in {"1", "L", "LA", "P", "RGB", "RGBA", "I", "I;16"}:
        image = image.convert("RGBA" if _alpha(image) else "RGB")
    kwargs = {}
    if quality is not None:
        kwargs["quality"] = quality
    if fmt == "png":
        kwargs["optimize"] = True
    if fmt == "tiff":
        kwargs["compression"] = "tiff_lzw"
    if metadata:
        for key in ("exif", "icc_profile", "dpi"):
            if metadata.get(key) and fmt not in {"gif", "bmp"}:
                kwargs[key] = metadata[key]
    Image.init()
    if fmt == "avif" and "AVIF" not in Image.SAVE:
        with tempfile.TemporaryDirectory(prefix="FormatWheel-") as tmp:
            image.save(Path(tmp) / "input.png")
            _command([_require("ffmpeg"), "-hide_banner", "-nostdin", "-y", "-loglevel", "error", "-i", str(Path(tmp) / "input.png"), "-frames:v", "1", "-vf", "pad=ceil(iw/2)*2:ceil(ih/2)*2,format=yuv420p", "-c:v", "libsvtav1", "-crf", str(round(52 - (quality or 62) * .34)), "-preset", "6", "-f", "avif", str(output)])
    else:
        image.save(output, format=formats[fmt], **kwargs)


def _same_image(source: Path) -> str:
    ext = _canonical(source.suffix)
    return ext if ext in {"jpg", "png", "heic", "tiff", "gif", "bmp", "webp", "avif"} else "png"


def _shrink(source: Path, output: Path) -> RunResult:
    before, after = source.stat().st_size, output.stat().st_size
    if after >= before:
        output.unlink(missing_ok=True)
        raise FormatWheelError("This file is already small; compression would make it larger.")
    return RunResult([output], f"Size reduced by {int((1 - after / max(1, before)) * 100)}%.")


def _image_action(action: str, source: Path) -> RunResult:
    if action == "img.to.pdf":
        return _merge([source], lambda _: None, suffix="")
    if action == "img.strip":
        return _strip_image(source)
    image, info = _load_image(source)
    if action == "img.ocr":
        text = _ocr(image)
        with _output(source, "txt", "Text") as out:
            out.write_text(text, encoding="utf-8")
        return RunResult([out], f"Recognized {len(text):,} characters; copied.", text)
    fmt, suffix, quality = _same_image(source), "", None
    quality = {"jpg": 90, "heic": 82, "webp": 84, "avif": 62}.get(fmt)
    if action.startswith("img.to."):
        fmt = action.rsplit(".", 1)[1]
        quality = {"jpg": 90, "heic": 82, "webp": 84, "avif": 62}.get(fmt)
    elif action == "img.half":
        suffix = "Half"
        image = image.resize((max(1, math.floor(image.width * .5 + .5)), max(1, math.floor(image.height * .5 + .5))), Image.Resampling.LANCZOS)
    elif action == "img.rotate":
        suffix = "Rotated"
        image = image.transpose(Image.Transpose.ROTATE_270)
        if fmt == "jpg":
            quality = 95
    elif action == "img.mono":
        suffix = "Monochrome"
        alpha = image.convert("RGBA").getchannel("A") if _alpha(image) else None
        image = ImageOps.grayscale(image).convert("RGB")
        if alpha is not None:
            image.putalpha(alpha)
    elif action == "img.compress":
        suffix, info = "Compressed", {}
        image.thumbnail((4096, 4096), Image.Resampling.LANCZOS)
        if fmt in {"jpg", "heic", "webp", "avif"}:
            quality = {"jpg": 70, "heic": 55, "webp": 70, "avif": 45}[fmt]
        else:
            fmt, quality = ("webp", 78) if _alpha(image) else ("jpg", 78)
    elif action == "img.cutout":
        image = _cutout(image)
        fmt, suffix, info = "png", "Cutout", {}
    elif action == "img.backdrop":
        image = _backdrop(image)
        fmt, suffix, info = "png", "Backdrop", {}
    else:
        raise FormatWheelError(f"Unknown action: {action}")
    with _output(source, fmt, suffix) as out:
        _save_image(image, out, fmt, quality, info)
        if action == "img.compress":
            return _shrink(source, out)
    note = f"{image.width}×{image.height}" if action == "img.half" else ""
    return RunResult([out], note)


def _backdrop(image: Image.Image) -> Image.Image:
    image.thumbnail((4000, 4000), Image.Resampling.LANCZOS)
    pad = max(1, round(max(image.size) * .09))
    width, height = image.width + pad * 2, image.height + pad * 2
    colors = [(250, 89, 128), (252, 140, 92), (252, 201, 102)]
    # A tiny diagonal ramp is smoothly enlarged, avoiding millions of Python iterations.
    ramp = Image.new("RGB", (256, 256))
    pixels = []
    for y in range(256):
        for x in range(256):
            t = (x + y) / 510
            a, b, f = (colors[0], colors[1], t / .55) if t <= .55 else (colors[1], colors[2], (t - .55) / .45)
            pixels.append(tuple(round(c + (d - c) * f) for c, d in zip(a, b)))
    ramp.putdata(pixels)
    canvas = ramp.resize((width, height), Image.Resampling.BILINEAR).convert("RGBA")
    mask = Image.new("L", image.size)
    ImageDraw.Draw(mask).rounded_rectangle((0, 0, image.width - 1, image.height - 1), radius=min(image.size) * .035, fill=255)
    shadow_mask = Image.new("L", canvas.size)
    shadow_mask.paste(mask, (pad, pad + round(pad * .12)))
    shadow_mask = shadow_mask.filter(ImageFilter.GaussianBlur(pad * .225)).point(lambda v: round(v * .35))
    canvas.alpha_composite(Image.composite(Image.new("RGBA", canvas.size, (0, 0, 0, 255)), Image.new("RGBA", canvas.size), shadow_mask))
    card = _flatten(image).convert("RGBA")
    card.putalpha(mask)
    canvas.alpha_composite(card, (pad, pad))
    return canvas.convert("RGB")


def _strip_image(source: Path) -> RunResult:
    fmt = _same_image(source)
    with _output(source, fmt, "Private") as out:
        if fmt == "jpg":
            data = source.read_bytes()
            if data[:2] != b"\xff\xd8":
                raise FormatWheelError("Invalid JPEG.")
            with Image.open(source) as image:
                orientation = image.getexif().get(274, 1)
            cleaned = bytearray(data[:2])
            if orientation != 1:
                exif = Image.Exif()
                exif[274] = orientation
                payload = exif.tobytes()
                cleaned += b"\xff\xe1" + struct.pack(">H", len(payload) + 2) + payload
            offset = 2
            while offset < len(data):
                start = offset
                if data[offset] != 255:
                    raise FormatWheelError("Invalid JPEG marker.")
                while offset < len(data) and data[offset] == 255:
                    offset += 1
                marker = data[offset]
                offset += 1
                if marker == 0xDA:  # Image entropy stream remains byte-for-byte identical.
                    cleaned += data[start:]
                    break
                if marker == 0xD9:
                    cleaned += data[start:offset]
                    break
                length = struct.unpack(">H", data[offset:offset + 2])[0]
                if length < 2 or offset + length > len(data):
                    raise FormatWheelError("Truncated JPEG.")
                payload = data[offset + 2:offset + length]
                color_marker = marker == 0xEE and payload.startswith(b"Adobe") or marker == 0xE2 and payload.startswith(b"ICC_PROFILE\x00")
                if color_marker or not (0xE1 <= marker <= 0xEF or marker == 0xFE):
                    cleaned += data[start:offset + length]
                offset += length
            out.write_bytes(cleaned)
        elif fmt == "png":
            data = source.read_bytes()
            if data[:8] != b"\x89PNG\r\n\x1a\n":
                raise FormatWheelError("Invalid PNG.")
            cleaned, offset = bytearray(data[:8]), 8
            with Image.open(source) as image:
                orientation = image.getexif().get(274, 1)
            # Preserve lossless pixel chunks, transparency, color management and animation.
            keep = {b"IHDR", b"PLTE", b"IDAT", b"IEND", b"tRNS", b"gAMA", b"cHRM", b"sRGB", b"acTL", b"fcTL", b"fdAT"}
            while offset + 12 <= len(data):
                length = struct.unpack(">I", data[offset:offset + 4])[0]
                end = offset + 12 + length
                if end > len(data):
                    raise FormatWheelError("Truncated PNG.")
                if data[offset + 4:offset + 8] in keep:
                    cleaned += data[offset:end]
                if data[offset + 4:offset + 8] == b"IHDR" and orientation != 1:
                    exif = Image.Exif()
                    exif[274] = orientation
                    payload = b"eXIf" + exif.tobytes()[6:]
                    cleaned += struct.pack(">I", len(payload) - 4) + payload + struct.pack(">I", zlib.crc32(payload) & 0xffffffff)
                offset = end
            out.write_bytes(cleaned)
        else:
            image, _ = _load_image(source)
            _save_image(image, out, fmt, 95 if fmt == "jpg" else None)
    return RunResult([out], "Removed location, device and author metadata.")


def _ocr(image: Image.Image) -> str:
    executable = _require("tesseract")
    tessdata = Path(os.environ.get("LOCALAPPDATA", str(Path.home() / "AppData/Local"))) / "FormatWheel/tessdata"
    data_args = ["--tessdata-dir", str(tessdata)] if tessdata.is_dir() else []
    with tempfile.TemporaryDirectory(prefix="FormatWheel-OCR-") as tmp:
        source = Path(tmp) / "page.png"
        image.save(source)
        installed = _command([executable] + data_args + ["--list-langs"]).splitlines()
        languages = [lang for lang in ("chi_sim", "chi_tra", "eng") if lang in installed]
        if not languages:
            raise FormatWheelError("Install Tesseract English / Chinese language data.")
        text = _command([executable, str(source), "stdout"] + data_args + ["-l", "+".join(languages), "--psm", "3"]).strip()
    if not text:
        raise FormatWheelError("No text was recognized.")
    return text


def _image_pdf_page(doc, image: Image.Image):
    fitz = _fitz()
    factor = 842 / max(image.size)
    page = doc.new_page(width=image.width * factor, height=image.height * factor)
    buffer = BytesIO()
    if _alpha(image):
        image.save(buffer, "PNG")
    else:
        _flatten(image).save(buffer, "JPEG", quality=90)
    page.insert_image(page.rect, stream=buffer.getvalue())


def _merge(paths: list[Path], progress: Progress, suffix: str = "Merged") -> RunResult:
    with _fitz().open() as merged:
        for index, path in enumerate(paths):
            if _kind(path) == "pdf":
                with _open_pdf(path) as source:
                    merged.insert_pdf(source)
            elif _kind(path) == "image":
                _image_pdf_page(merged, _load_image(path)[0])
            progress((index + 1) / len(paths))
        if not len(merged):
            raise FormatWheelError("No PDF pages or images to merge.")
        merged.set_metadata({"creator": "FormatWheel"})
        with _output(paths[0], "pdf", suffix) as out:
            _save_pdf(merged, out, garbage=4, deflate=True)
    return RunResult([out])


def _render(page, scale: float = 2) -> Image.Image:
    fitz = _fitz()
    area = page.rect.width * page.rect.height
    scale = min(scale, math.sqrt(60_000_000 / max(1, area)))
    pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), colorspace=fitz.csRGB, alpha=False)
    return Image.frombytes("RGB", (pix.width, pix.height), pix.samples)


def _pdf_action(action: str, source: Path, progress: Progress) -> RunResult:
    if action in {"pdf.to.docx", "pdf.to.txt", "pdf.to.md", "pdf.to.rtf", "pdf.to.html"}:
        paragraphs = _pdf_paragraphs(source)
        with _output(source, action.rsplit(".", 1)[1]) as out:
            _write_paragraphs(paragraphs, out, action.rsplit(".", 1)[1])
        return RunResult([out])
    fitz = _fitz()
    with _open_pdf(source) as doc:
        count = len(doc)
        if action in {"pdf.to.png", "pdf.to.jpg"}:
            fmt = action.rsplit(".", 1)[1]
            if count == 1:
                with _output(source, fmt) as out:
                    _save_image(_render(doc[0]), out, fmt, 90 if fmt == "jpg" else None)
            else:
                with _folder(source, fmt.upper()) as out:
                    digits = max(2, len(str(count)))
                    for index, page in enumerate(doc):
                        _save_image(_render(page), out / f"{source.stem}-{index + 1:0{digits}d}.{fmt}", fmt, 90 if fmt == "jpg" else None)
                        progress((index + 1) / count)
            return RunResult([out], f"{count} page(s).")
        if action == "pdf.to.long":
            return _long_pdf(doc, source, progress)
        if action == "pdf.ocr":
            parts = []
            for index, page in enumerate(doc):
                try:
                    parts.append(_ocr(_render(page)))
                except FormatWheelError as exc:
                    if str(exc) != "No text was recognized.":
                        raise
                progress((index + 1) / count)
            text = "\n\n".join(parts).strip()
            if not text:
                raise FormatWheelError("No text was recognized.")
            with _output(source, "txt", "Text") as out:
                out.write_text(text, encoding="utf-8")
            return RunResult([out], f"Recognized {len(text):,} characters; copied.", text)
        if action == "pdf.split":
            if count < 2:
                raise FormatWheelError("This PDF has only one page; there is nothing to split.")
            with _folder(source, "Split") as out:
                for index in range(count):
                    with fitz.open() as single:
                        single.insert_pdf(doc, from_page=index, to_page=index)
                        _save_pdf(single, out / f"{source.stem}-{index + 1}.pdf", garbage=4, deflate=True)
                    progress((index + 1) / count)
            return RunResult([out], f"Split into {count} files.")
        suffix = {"pdf.rotate": "Rotated", "pdf.mono": "Monochrome", "pdf.compress": "Compressed", "pdf.strip": "Private"}.get(action)
        if suffix is None:
            raise FormatWheelError(f"Unknown action: {action}")
        with _output(source, "pdf", suffix) as out:
            if action == "pdf.rotate":
                for page in doc:
                    page.set_rotation((page.rotation + 90) % 360)
            elif action == "pdf.mono":
                if not hasattr(doc[0], "recolor"):
                    raise FormatWheelError("Upgrade PyMuPDF to enable vector-preserving PDF grayscale conversion.")
                for index, page in enumerate(doc):
                    page.recolor(components=1)
                    progress((index + 1) / count)
            elif action == "pdf.strip":
                metadata = doc.metadata or {}
                for key in ("author", "creator", "producer", "subject", "keywords", "creationDate", "modDate"):
                    metadata[key] = ""
                doc.set_metadata(metadata)
                doc.del_xml_metadata()
            elif action == "pdf.compress" and hasattr(doc, "rewrite_images"):
                doc.rewrite_images(dpi_threshold=150, dpi_target=120, quality=70)
            _save_pdf(doc, out, garbage=4, deflate=True, clean=True)
            if action == "pdf.compress":
                return _shrink(source, out)
    return RunResult([out], "Removed author, software and date metadata." if action == "pdf.strip" else "")


def _long_pdf(doc, source: Path, progress: Progress) -> RunResult:
    count = min(len(doc), 120)
    outputs, images, height = [], [], 0
    def flush():
        nonlocal images, height
        if not images:
            return
        canvas = Image.new("RGB", (1242, height - 16), (230, 230, 230))
        y = 0
        for image in images:
            canvas.paste(image, (0, y))
            y += image.height + 16
        with _output(source, "jpg", f"Long {len(outputs) + 1}") as out:
            _save_image(canvas, out, "jpg", 88)
        outputs.append(out)
        images, height = [], 0
    try:
        for index in range(count):
            page = doc[index]
            image = _render(page, 1242 / max(1, page.rect.width))
            if image.width != 1242:
                image = image.resize((1242, max(1, round(image.height * 1242 / image.width))), Image.Resampling.LANCZOS)
            # A single unusually tall page is split as well, respecting JPEG limits.
            for top in range(0, image.height, 59000):
                piece = image.crop((0, top, image.width, min(top + 59000, image.height)))
                if images and height + piece.height > 60000:
                    flush()
                images.append(piece)
                height += piece.height + 16
            progress((index + 1) / count * .9)
        flush()
    except BaseException:
        for output in outputs:
            output.unlink(missing_ok=True)
        raise
    return RunResult(outputs, f"{count} pages stitched." + (f" First 120 of {len(doc)} pages." if len(doc) > 120 else ""))


def _probe(source: Path) -> dict:
    return json.loads(_command([_require("ffprobe"), "-v", "error", "-show_format", "-show_streams", "-of", "json", str(source)]))


def _ff(source: Path, arguments: list[str], output: Path, progress: Progress, duration: float | None = None):
    if duration is None:
        try:
            duration = float(_probe(source).get("format", {}).get("duration", 0))
        except (FormatWheelError, ValueError):
            duration = 0
    args = [_require("ffmpeg"), "-hide_banner", "-nostdin", "-y", "-loglevel", "error", "-progress", "pipe:1", "-nostats", "-i", str(source)] + arguments + [str(output)]
    # stderr goes to a temporary file so progress reads can never deadlock on a full pipe.
    with tempfile.TemporaryFile() as errors:
        process = subprocess.Popen(args, stdin=subprocess.DEVNULL, stdout=subprocess.PIPE, stderr=errors,
                                   text=True, encoding="utf-8", errors="replace",
                                   creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        try:
            for line in process.stdout:
                if duration and line.startswith("out_time_us="):
                    try:
                        progress(min(1, max(0, float(line.split("=", 1)[1]) / 1_000_000 / duration)))
                    except ValueError:
                        pass
            code = process.wait()
        except BaseException:
            process.kill()
            process.wait()
            raise
        finally:
            process.stdout.close()
        if code:
            errors.seek(0)
            lines = errors.read().decode("utf-8", "replace").strip().splitlines()
            raise FormatWheelError("FFmpeg: " + (lines[-1][:350] if lines else f"exit code {code}"))
    progress(1)


def _scale(side: int) -> str:
    return f"scale='if(gt(iw,ih),min({side},iw),-2)':'if(gt(iw,ih),-2,min({side},ih))':flags=lanczos"


def _media_action(action: str, source: Path, progress: Progress) -> RunResult:
    probe = _probe(source)
    streams = probe.get("streams", [])
    duration = float(probe.get("format", {}).get("duration", 0) or 0)
    audio = next((s for s in streams if s.get("codec_type") == "audio"), None)
    video = next((s for s in streams if s.get("codec_type") == "video" and not s.get("disposition", {}).get("attached_pic")), None)
    h264 = ["-c:v", "libx264", "-preset", "medium", "-crf", "20", "-pix_fmt", "yuv420p", "-c:a", "aac", "-b:a", "192k"]
    ext, suffix, args, note = "mp4", "", [], ""
    effective_duration = duration
    if action.startswith("vid."):
        if not video:
            raise FormatWheelError("This file has no video stream.")
        if action in {"vid.to.mp4", "vid.to.mov"}:
            ext = action.rsplit(".", 1)[1]
            # Use passthrough where the codecs are compatible with the destination.
            compatible = video.get("codec_name") in {"h264", "hevc", "mpeg4"} and (not audio or audio.get("codec_name") in {"aac", "mp3", "alac"})
            args = ["-map", "0:v:0", "-map", "0:a?", "-c", "copy"] if compatible else h264
            if ext == "mp4":
                args += ["-movflags", "+faststart"]
        elif action == "vid.to.hevc":
            suffix = "HEVC"
            args = ["-c:v", "libx265", "-preset", "medium", "-crf", "24", "-pix_fmt", "yuv420p", "-tag:v", "hvc1", "-c:a", "aac", "-b:a", "192k", "-movflags", "+faststart"]
        elif action == "vid.to.webm":
            ext = "webm"
            args = ["-c:v", "libvpx-vp9", "-crf", "32", "-b:v", "0", "-row-mt", "1", "-deadline", "good", "-cpu-used", "4", "-c:a", "libopus", "-b:a", "128k"]
        elif action == "vid.to.gif":
            ext = "gif"
            graph = f"fps=12,{_scale(540)},split[a][b];[a]palettegen=max_colors=200:stats_mode=diff[p];[b][p]paletteuse=dither=bayer:bayer_scale=4:diff_mode=rectangle"
            args = ["-t", "30", "-filter_complex", graph, "-loop", "0", "-an"]
            effective_duration = min(duration, 30)
            note = "Only the first 30 seconds." if duration > 30 else ""
        elif action in {"vid.to.mp3", "vid.to.m4a", "vid.to.wav"}:
            if not audio:
                raise FormatWheelError("This video has no audio stream.")
            ext = action.rsplit(".", 1)[1]
            args = ["-vn"] + {"mp3": ["-c:a", "libmp3lame", "-q:a", "2"], "m4a": ["-c:a", "aac", "-b:a", "192k"], "wav": ["-c:a", "pcm_s16le"]}[ext]
        elif action == "vid.compress":
            suffix = "Compressed"
            args = ["-c:v", "libx264", "-preset", "faster", "-crf", "27", "-vf", _scale(1920), "-pix_fmt", "yuv420p", "-c:a", "aac", "-b:a", "128k", "-movflags", "+faststart"]
        elif action == "vid.mute":
            suffix = "Muted"
            ext = source.suffix.lower().lstrip(".") if source.suffix.lower() in {".mp4", ".mov", ".m4v", ".mkv", ".webm"} else "mp4"
            args = ["-map", "0:v", "-c", "copy", "-an"]
        elif action == "vid.720p":
            suffix = "720P"
            args = ["-vf", _scale(1280)] + h264 + ["-movflags", "+faststart"]
        elif action == "vid.cover":
            ext, suffix = "jpg", "Cover"
            at = min(1, duration * .1) if duration > 2 else 0
            args = ["-ss", str(at), "-frames:v", "1", "-q:v", "2", "-an"]
        elif action in {"vid.fast", "vid.slow"}:
            fast = action == "vid.fast"
            suffix = "2x" if fast else "Slow"
            args = ["-filter:v", "setpts=0.5*PTS" if fast else "setpts=2.0*PTS"]
            if audio:
                args += ["-filter:a", "atempo=2.0" if fast else "atempo=0.5"]
            args += ["-c:v", "libx264", "-preset", "fast", "-crf", "20", "-pix_fmt", "yuv420p", "-c:a", "aac", "-b:a", "160k", "-movflags", "+faststart"]
            effective_duration = duration * (.5 if fast else 2)
        elif action == "vid.strip":
            suffix, ext = "Private", source.suffix.lower().lstrip(".") or "mp4"
            args = ["-map", "0", "-map_metadata", "-1", "-map_chapters", "-1", "-c", "copy"]
            for index in range(len(streams)):
                args += [f"-map_metadata:s:{index}", "-1"]
            note = "Removed location, device and capture metadata."
        else:
            raise FormatWheelError(f"Unknown action: {action}")
    else:
        if not audio:
            raise FormatWheelError("This file has no audio stream.")
        codecs = {"mp3": ["-c:a", "libmp3lame", "-q:a", "2"], "m4a": ["-c:a", "aac", "-b:a", "256k"], "wav": ["-c:a", "pcm_s16le"], "aiff": ["-c:a", "pcm_s16be"], "flac": ["-c:a", "flac"], "opus": ["-c:a", "libopus", "-b:a", "128k"], "alac": ["-c:a", "alac"], "m4r": ["-t", "40", "-c:a", "aac", "-b:a", "256k", "-f", "ipod"]}
        if action.startswith("aud.to."):
            target = action.rsplit(".", 1)[1]
            if target not in codecs:
                raise FormatWheelError(f"Unknown action: {action}")
            ext = "m4a" if target == "alac" else target
            suffix = "Lossless" if target == "alac" else "Ringtone" if target == "m4r" else ""
            args = ["-vn"] + codecs[target]
            if target == "m4r":
                effective_duration = min(duration, 40)
                note = "Only the first 40 seconds." if duration > 40 else ""
        elif action == "aud.compress":
            ext, suffix = "m4a", "Compressed"
            args = ["-vn", "-c:a", "aac", "-b:a", "96k"]
        elif action in {"aud.mono", "aud.loud", "aud.fast"}:
            ext = _canonical(source.suffix)
            if ext not in {"mp3", "wav", "aiff", "flac", "opus"}:
                ext = "m4a"
            suffix = {"aud.mono": "Mono", "aud.loud": "Normalized", "aud.fast": "1.5x"}[action]
            rate = str(audio.get("sample_rate", 48000))
            filter_args = {"aud.mono": ["-ac", "1"], "aud.loud": ["-af", "loudnorm=I=-16:TP=-1.5:LRA=11", "-ar", rate], "aud.fast": ["-af", "atempo=1.5"]}[action]
            args = ["-vn"] + filter_args + codecs[ext]
            if action == "aud.fast":
                effective_duration = duration / 1.5
        else:
            raise FormatWheelError(f"Unknown action: {action}")
    with _output(source, ext, suffix) as out:
        _ff(source, args, out, progress, effective_duration)
        if action.endswith(".compress"):
            return _shrink(source, out)
    return RunResult([out], note)


@dataclass
class _Paragraph:
    text: str
    size: float = 12
    bold: bool = False
    heading: int = 0
    runs: list[tuple[str, bool]] = field(default_factory=list)


_SUPPLEMENT = dict(zip(
    [0x2E9F, 0x2EA0, 0x2EC4, 0x2EC5, 0x2EC6, 0x2EC8, 0x2EC9, 0x2ECB, 0x2ECE, 0x2ED0, 0x2ED3, 0x2ED4, 0x2ED7, 0x2ED8, 0x2EDA, 0x2EDB, 0x2EDC, 0x2EDD, 0x2EE0, 0x2EE2, 0x2EE4, 0x2EE5, 0x2EE6, 0x2EE7, 0x2EE8, 0x2EE9, 0x2EEA, 0x2EEC, 0x2EEE, 0x2EF0, 0x2EF3],
    "母民西见角讠贝车辶钅长门雨青页风飞食饣马鬼鱼鸟卤麦黄黾齐齿龙龟"))


def _text_fix(text: str) -> str:
    return "".join(unicodedata.normalize("NFKC", c) if 0x2F00 <= ord(c) <= 0x2FD5 else _SUPPLEMENT.get(ord(c), c) for c in text)


def _cjk(c: str) -> bool:
    return bool(c) and (0x2E80 <= ord(c) <= 0x9FFF or 0xF900 <= ord(c) <= 0xFAFF or 0xFF00 <= ord(c) <= 0xFFEF or 0x20000 <= ord(c) <= 0x2FA1F)


def _pdf_paragraphs(source: Path) -> list[_Paragraph]:
    paragraphs = []
    with _open_pdf(source) as doc:
        for page in doc:
            lines = []
            for block in page.get_text("dict", sort=True)["blocks"]:
                for line in block.get("lines", []):
                    spans = line.get("spans", [])
                    text = _text_fix("".join(s["text"] for s in spans).strip())
                    if text:
                        runs = [(_text_fix(s["text"]), bool(s["flags"] & 16)) for s in spans]
                        if runs:
                            runs[0] = (runs[0][0].lstrip(), runs[0][1])
                            runs[-1] = (runs[-1][0].rstrip(), runs[-1][1])
                        lines.append((text, max(s["size"] for s in spans), line["bbox"], all(s["flags"] & 16 for s in spans if s["text"].strip()), runs))
            if not lines:
                continue
            gaps = sorted(b[2][1] - a[2][3] for a, b in zip(lines, lines[1:]) if abs(a[1] - b[1]) < .6 and b[2][1] - a[2][3] > -1)
            usual = gaps[len(gaps) // 2] if gaps else 0
            widest = max(line[2][2] - line[2][0] for line in lines)
            previous = None
            current = None
            for text, size, box, bold, runs in lines:
                split = previous is None
                if previous:
                    old_text, old_size, old_box, _, _ = previous
                    gap = box[1] - old_box[3]
                    height = max(box[3] - box[1], old_box[3] - old_box[1])
                    split = (abs(size - old_size) > .6 or gap > usual + height * .35 or gap < -height * 1.5 or bool(re.match(r"^(?:[•◦▪●·‣–—*]|[-]\s|\d{1,3}[.)、]|[a-zA-Z][.)、]|[（(]\d{1,3}[)）]|[一二三四五六七八九十]+、)", text)) or ((old_box[2] - old_box[0]) < widest * .72 and old_text[-1:] in "。！？!?.:：;；」”)）"))
                if split:
                    current = _Paragraph(text, size, bold, runs=runs)
                    paragraphs.append(current)
                else:
                    separator = "" if _cjk(current.text[-1:]) or _cjk(text[:1]) else " "
                    current.text += separator + text
                    current.bold &= bold
                    current.runs += ([(separator, False)] if separator else []) + runs
                previous = (text, size, box, bold, runs)
            paragraphs.append(_Paragraph(""))
    if not any(p.text.strip() for p in paragraphs):
        raise FormatWheelError("This PDF has no text layer. Use OCR in Tools instead.")
    sizes = Counter(round(p.size) for p in paragraphs for _ in range(min(10000, len(p.text))))
    body = sizes.most_common(1)[0][0] if sizes else 12
    for p in paragraphs:
        p.heading = 1 if p.size >= body * 1.55 else 2 if p.size >= body * 1.28 else 3 if p.size >= body * 1.1 or p.bold and len(p.text) <= 40 else 0
    return paragraphs


def _paragraph_html(paragraphs: list[_Paragraph]) -> str:
    blocks = []
    for p in paragraphs:
        tag = f"h{p.heading}" if p.heading else "p"
        if p.runs and not p.heading:
            text = "".join(("<b>" + escape(t) + "</b>" if bold and t.strip() else escape(t)) for t, bold in p.runs).replace("\n", "<br>")
        else:
            text = escape(p.text).replace("\n", "<br>")
        if p.bold and not p.heading and not p.runs:
            text = "<b>" + text + "</b>"
        blocks.append(f"<{tag}>{text or '&nbsp;'}</{tag}>")
    return '<!doctype html><html><head><meta charset="utf-8"><style>body{font-family:Calibri,"Microsoft YaHei",sans-serif;font-size:12pt;line-height:1.5}p{margin:0 0 8pt}pre,code{font-family:Consolas,monospace}</style></head><body>' + "\n".join(blocks) + "</body></html>"


def _rtf_escape(text: str) -> str:
    pieces = []
    for char in text:
        if char in "\\{}":
            pieces.append("\\" + char)
        elif char == "\n":
            pieces.append("\\line ")
        elif 32 <= ord(char) < 127:
            pieces.append(char)
        else:
            for n in struct.unpack("<" + "H" * (len(char.encode("utf-16le")) // 2), char.encode("utf-16le")):
                pieces.append(f"\\u{n if n < 32768 else n - 65536}?")
    return "".join(pieces)


def _write_paragraphs(paragraphs: list[_Paragraph], output: Path, fmt: str):
    if fmt == "txt":
        output.write_text("\n".join(p.text for p in paragraphs).strip() + "\n", encoding="utf-8")
    elif fmt == "md":
        def inline(p):
            if p.heading or not p.runs:
                return p.text
            pieces = []
            for text, bold in p.runs:
                trimmed = text.strip()
                if bold and trimmed:
                    lead = text[:len(text) - len(text.lstrip())]
                    tail = text[len(text.rstrip()):]
                    pieces.append(lead + "**" + trimmed + "**" + tail)
                else:
                    pieces.append(text)
            return "".join(pieces)
        output.write_text("\n\n".join(("#" * p.heading + " " if p.heading else "") + inline(p) for p in paragraphs if p.text).strip() + "\n", encoding="utf-8")
    elif fmt == "html":
        output.write_text(_paragraph_html(paragraphs), encoding="utf-8")
    elif fmt == "rtf":
        text = r"{\rtf1\ansi\deff0{\fonttbl{\f0 Calibri;}}\uc1 "
        for p in paragraphs:
            text += f"\\pard\\f0\\fs{round(p.size * 2)}"
            for content, bold in (p.runs or [(p.text, p.bold)]):
                text += ("\\b " if bold or p.heading else "\\b0 ") + _rtf_escape(content)
            text += "\\par\n"
        output.write_text(text + "}", encoding="ascii")
    elif fmt == "docx":
        from docx import Document
        from docx.shared import Pt
        doc = Document()
        for p in paragraphs:
            paragraph = doc.add_paragraph(style=f"Heading {p.heading}" if p.heading else None)
            for content, bold in (p.runs or [(p.text, p.bold)]):
                run = paragraph.add_run(content)
                run.bold, run.font.size = bold or bool(p.heading), Pt(p.size)
        doc.save(output)
    elif fmt == "pdf":
        _html_pdf(_paragraph_html(paragraphs), output)
    else:
        raise FormatWheelError(f"Unsupported document format: {fmt}")


def _html_pdf(html: str, output: Path):
    fitz = _fitz()
    story = fitz.Story(html=html)
    with fitz.DocumentWriter(str(output)) as writer:
        more = True
        while more:
            device = writer.begin_page(fitz.Rect(0, 0, 595, 842))
            more, _ = story.place(fitz.Rect(60, 56, 535, 786))
            story.draw(device)
            writer.end_page()
    _stamp_pdf_file(output)


def _read_text(source: Path) -> str:
    data = source.read_bytes()
    for encoding in ("utf-8-sig", "utf-16" if data[:2] in {b"\xff\xfe", b"\xfe\xff"} else "utf-8", "gb18030"):
        try:
            return data.decode(encoding).replace("\r\n", "\n").replace("\r", "\n")
        except UnicodeError:
            continue
    return data.decode("cp1252", errors="replace").replace("\r\n", "\n").replace("\r", "\n")


def _markdown_plain(text: str) -> str:
    # Match the upstream intentionally small Markdown-to-plain transformation.
    text = re.sub(r"(?m)^\s*#{1,6} ", "", text)
    return text.replace("**", "").replace("__", "").replace("`", "")


def _docx_paragraphs(source: Path) -> list[_Paragraph]:
    from docx import Document
    from docx.text.paragraph import Paragraph
    from docx.table import Table
    doc = Document(source)
    paragraphs = []
    for child in doc.element.body:
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, doc)
            heading = re.match(r"Heading ([1-6])", p.style.name or "")
            sizes = [run.font.size.pt for run in p.runs if run.font.size]
            texts = [run for run in p.runs if run.text.strip()]
            paragraphs.append(_Paragraph(p.text, max(sizes, default=12), bool(texts) and all(run.bold for run in texts), int(heading[1]) if heading else 0, [(run.text, bool(run.bold)) for run in p.runs]))
        elif tag == "tbl":
            table = Table(child, doc)
            paragraphs.extend(_Paragraph("\t".join(cell.text for cell in row.cells)) for row in table.rows)
    return paragraphs


class _HTMLText(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts, self.hidden = [], 0
    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style", "head"}:
            self.hidden += 1
        elif not self.hidden and tag in {"p", "div", "br", "li", "h1", "h2", "h3", "h4", "tr"}:
            self.parts.append("\n")
    def handle_endtag(self, tag):
        if tag in {"script", "style", "head"}:
            self.hidden = max(0, self.hidden - 1)
        elif not self.hidden and tag in {"p", "div", "li", "h1", "h2", "h3", "h4", "tr"}:
            self.parts.append("\n")
    def handle_data(self, data):
        if not self.hidden:
            self.parts.append(data)


def _office_convert(source: Path, target: Path, fmt: str):
    filters = {"pdf": "pdf:writer_pdf_Export", "docx": "docx:Office Open XML Text", "doc": "doc:MS Word 97", "rtf": "rtf:Rich Text Format", "odt": "odt:writer8", "html": "html:HTML (StarWriter)", "txt": "txt:Text (encoded):UTF8"}
    if fmt not in filters:
        raise FormatWheelError(f"Unsupported LibreOffice output: {fmt}")
    with tempfile.TemporaryDirectory(prefix="FormatWheel-Office-") as tmp:
        folder = Path(tmp)
        incoming = source
        if source.suffix.lower() == ".rtfd":
            if not source.is_dir():
                raise FormatWheelError("This RTFD file is not an unpacked document package.")
            incoming = next((p for p in source.iterdir() if p.name.lower() == "txt.rtf"), None)
            if incoming is None:
                raise FormatWheelError("This RTFD package has no TXT.rtf document.")
        elif source.suffix.lower() == ".wordml":
            incoming = folder / "input.xml"
            shutil.copyfile(source, incoming)
        elif source.suffix.lower() == ".webarchive":
            try:
                archive = plistlib.loads(source.read_bytes())
                main = archive["WebMainResource"]
                html = main["WebResourceData"].decode(main.get("WebResourceTextEncodingName", "utf-8"), errors="replace")
                import base64
                for resource in archive.get("WebSubresources", []):
                    url, payload = resource.get("WebResourceURL", ""), resource.get("WebResourceData", b"")
                    if url and payload:
                        html = html.replace(url, "data:" + resource.get("WebResourceMIMEType", "application/octet-stream") + ";base64," + base64.b64encode(payload).decode("ascii"))
                incoming = folder / "input.html"
                incoming.write_text(html, encoding="utf-8")
            except (KeyError, ValueError, TypeError, LookupError) as exc:
                raise FormatWheelError("This WebArchive has no readable embedded webpage.") from exc
        if _canonical(source.suffix) == "md":
            import markdown
            incoming = folder / "input.html"
            html = markdown.markdown(_read_text(source), extensions=["tables", "fenced_code"])
            # Upstream uses image alt text; it never fetches linked pictures.
            html = re.sub(r"<img\b[^>]*alt=\"([^\"]*)\"[^>]*>", r"\1", html)
            incoming.write_text('<html><head><meta charset="utf-8"></head><body>' + html + "</body></html>", encoding="utf-8")
        output_folder = folder / "out"
        output_folder.mkdir()
        profile = (folder / "profile").as_uri()
        _command([_require("soffice"), f"-env:UserInstallation={profile}", "--headless", "--nologo", "--nodefault", "--norestore", "--convert-to", filters[fmt], "--outdir", str(output_folder), str(incoming)])
        candidates = [p for p in output_folder.iterdir() if p.suffix.lower() == "." + fmt]
        if not candidates:
            raise FormatWheelError("LibreOffice could not read this document or produce the requested format.")
        shutil.copyfile(candidates[0], target)
        if fmt == "pdf":
            _stamp_pdf_file(target)
        # HTML exports can include pictures in a sibling directory; keep their URLs valid.
        if fmt == "html":
            content = target.read_text(encoding="utf-8", errors="replace")
            for asset in output_folder.iterdir():
                if asset == candidates[0]:
                    continue
                if asset.is_file():
                    mime = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".gif": "image/gif", ".svg": "image/svg+xml"}.get(asset.suffix.lower())
                    if mime:
                        import base64
                        content = content.replace(asset.name, "data:" + mime + ";base64," + base64.b64encode(asset.read_bytes()).decode("ascii"))
            target.write_text(content, encoding="utf-8")


def _plain_document(source: Path) -> str:
    fmt = _canonical(source.suffix)
    if fmt == "txt":
        return _read_text(source)
    if fmt == "md":
        return _markdown_plain(_read_text(source))
    if fmt == "docx":
        return "\n".join(p.text for p in _docx_paragraphs(source))
    if fmt == "html":
        parser = _HTMLText()
        parser.feed(_read_text(source))
        return re.sub(r"\n{3,}", "\n\n", "".join(parser.parts)).strip()
    with tempfile.TemporaryDirectory(prefix="FormatWheel-Text-") as tmp:
        output = Path(tmp) / "text.txt"
        _office_convert(source, output, "txt")
        return _read_text(output)


def _doc_action(action: str, source: Path) -> RunResult:
    if action == "doc.strip":
        return _strip_docx(source)
    if action == "doc.copy":
        text = _plain_document(source).strip()
        if not text:
            raise FormatWheelError("This document has no text.")
        return RunResult(note=f"Copied {len(text):,} characters.", clip=text)
    if action == "doc.count":
        text = _plain_document(source)
        cjk = words = other = 0
        in_word = False
        for char in text:
            point = ord(char)
            if 0x4E00 <= point <= 0x9FFF or 0x3400 <= point <= 0x4DBF or 0x20000 <= point <= 0x2A6DF or 0xF900 <= point <= 0xFAFF:
                cjk += 1
                in_word = False
            else:
                if char.isalnum():
                    if not in_word:
                        words += 1
                    in_word = True
                else:
                    in_word = False
                if not char.isspace():
                    other += 1
        note = f"{cjk + words:,} words/Chinese characters"
        if cjk and words:
            note += f" · Chinese {cjk:,} · other words {words:,}"
        return RunResult(note=note + f" · {cjk + other:,} characters excluding whitespace.")
    if not action.startswith("doc.to."):
        raise FormatWheelError(f"Unknown action: {action}")
    target, fmt = action.rsplit(".", 1)[1], _canonical(source.suffix)
    with _output(source, target) as out:
        if target == fmt:
            shutil.copyfile(source, out)
        elif target == "txt":
            out.write_text(_plain_document(source), encoding="utf-8")
        elif target == "md":
            if fmt == "md":
                out.write_text(_read_text(source), encoding="utf-8")
            elif fmt == "docx":
                _write_paragraphs(_docx_paragraphs(source), out, "md")
            elif fmt == "txt":
                out.write_text(_read_text(source), encoding="utf-8")
            else:
                with tempfile.TemporaryDirectory(prefix="FormatWheel-Markdown-") as tmp:
                    intermediate = Path(tmp) / "document.docx"
                    _office_convert(source, intermediate, "docx")
                    _write_paragraphs(_docx_paragraphs(intermediate), out, "md")
        elif fmt == "txt" and target in {"docx", "html", "pdf"}:
            _write_paragraphs([_Paragraph(line) for line in _read_text(source).splitlines()], out, target)
        elif fmt == "docx" and target == "html" and not find_tool("soffice"):
            _write_paragraphs(_docx_paragraphs(source), out, target)
        else:
            _office_convert(source, out, target)
    return RunResult([out])


def _xml_value(xml: bytes, name: str) -> str:
    root = ET.fromstring(xml)
    for element in root.iter():
        if element.tag.rsplit("}", 1)[-1] == name:
            return (element.text or "").strip()
    return ""


def _strip_docx(source: Path) -> RunResult:
    if source.suffix.lower() != ".docx":
        raise FormatWheelError("Privacy stripping is available for .docx documents.")
    fields = {"docProps/core.xml": {"creator", "lastModifiedBy", "keywords", "description", "subject", "category"}, "docProps/app.xml": {"Company", "Manager", "Template"}}
    with _output(source, "docx", "Private") as out:
        with zipfile.ZipFile(source) as original, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as cleaned:
            for entry in original.infolist():
                payload = original.read(entry)
                if entry.filename in fields:
                    # Preserve package namespace prefixes and unknown XML parts exactly.
                    text = payload.decode("utf-8")
                    for name in fields[entry.filename]:
                        text = re.sub(rf"(<(?:[\w.-]+:)?{name}(?:\s[^>]*)?>)[^<]*(</(?:[\w.-]+:)?{name}\s*>)", r"\1\2", text)
                    payload = text.encode("utf-8")
                cleaned.writestr(entry, payload)
    return RunResult([out], "Removed author, editor and company metadata.")


def scan_privacy(path: Path | str | Iterable[Path | str]) -> Wanted:
    """Match the original wanted meter, using the worst result in a selection."""
    if not isinstance(path, (str, Path)):
        findings = [scan_privacy(p) for p in list(path)[:24]]
        worst = max(findings, key=lambda w: w.level, default=Wanted())
        flagged = sum(w.level > 0 for w in findings)
        return Wanted(worst.level, worst.reasons + ([f"{flagged} files contain metadata"] if flagged > 1 else []))
    source, wanted = Path(path), Wanted()
    try:
        kind = _kind(source)
        if kind == "image":
            with Image.open(source) as image:
                exif = image.getexif()
                try:
                    gps = exif.get_ifd(34853)
                except (KeyError, ValueError, TypeError):
                    gps = {}
                try:
                    nested = exif.get_ifd(34665)
                except (KeyError, ValueError, TypeError):
                    nested = {}
                tags = dict(exif)
                tags.update(nested)
                if gps.get(2) is not None:
                    wanted.level += 2
                    wanted.reasons.append("GPS location")
                model = tags.get(272) or tags.get(42036)
                if model:
                    wanted.level += 1
                    wanted.reasons.append(str(model)[:24])
                if tags.get(36867):
                    wanted.level += 1
                    wanted.reasons.append("Capture date")
                if tags.get(42033) or tags.get(42032) or tags.get(315):
                    wanted.level += 1
                    wanted.reasons.append("Device serial / author")
        elif kind == "pdf":
            with _open_pdf(source) as doc:
                metadata = doc.metadata or {}
                if metadata.get("author", "").strip():
                    wanted.level += 2
                    wanted.reasons.append("Author: " + metadata["author"][:16])
                if metadata.get("creator", "").strip():
                    wanted.level += 1
                    wanted.reasons.append("Creation software")
                if metadata.get("keywords", "").strip() or metadata.get("subject", "").strip():
                    wanted.level += 1
                    wanted.reasons.append("Subject / keywords")
        elif kind == "document" and source.suffix.lower() == ".docx":
            with zipfile.ZipFile(source) as archive:
                core = archive.read("docProps/core.xml") if "docProps/core.xml" in archive.namelist() else b"<core/>"
                creator, editor = _xml_value(core, "creator"), _xml_value(core, "lastModifiedBy")
                if creator:
                    wanted.level += 2
                    wanted.reasons.append("Author: " + creator[:16])
                if editor and editor != creator:
                    wanted.level += 1
                    wanted.reasons.append("Editor: " + editor[:16])
                if "docProps/app.xml" in archive.namelist() and _xml_value(archive.read("docProps/app.xml"), "Company"):
                    wanted.level += 1
                    wanted.reasons.append("Company")
        elif kind == "video" and find_tool("ffprobe"):
            probe = _probe(source)
            tags = dict(probe.get("format", {}).get("tags", {}))
            for stream in probe.get("streams", []):
                tags.update(stream.get("tags", {}))
            tags = {key.lower(): value for key, value in tags.items()}
            if any("location" in key or key == "©xyz" for key in tags):
                wanted.level += 2
                wanted.reasons.append("GPS location")
            model = next((v for k, v in tags.items() if k.endswith("model")), "")
            if model:
                wanted.level += 1
                wanted.reasons.append(str(model)[:24])
            if "com.apple.quicktime.creationdate" in tags:
                wanted.level += 1
                wanted.reasons.append("Capture date")
    except Exception:
        # Scanning is advisory; unsupported or corrupt metadata must not block the wheel.
        return Wanted()
    wanted.level = min(wanted.level, 5)
    return wanted
